from pathlib import Path
import re
import pandas as pd
from docx import Document

DOCX_PATH = Path(__file__).resolve().parent.parent / "QA_pairs.docx"
OUT_CSV   = Path(__file__).resolve().parent.parent / "data" / "qa_pairs.csv"
OUT_CSV.parent.mkdir(parents=True, exist_ok=True)

# Matches like: "1. Question? Answer"
NUM_LINE_PAT = re.compile(r"^\s*(\d+)\s*[\.\)]\s*(.+?)\s*$")

def norm(s: str) -> str:
    return re.sub(r"\s+", " ", (s or "").strip())

def split_qa(text: str):
    """
    'When was RUSL established? November 1995.'
    -> ('When was RUSL established?', 'November 1995.')
    """
    t = norm(text)
    if "?" not in t:
        return None, None
    q, a = t.split("?", 1)
    q = norm(q) + "?"
    a = norm(a)
    return q if q else None, a

def collect_all_text(doc: Document):
    """
    Returns list of lines from:
    - paragraphs
    - table cells
    """
    lines = []

    # paragraphs
    for p in doc.paragraphs:
        t = norm(p.text)
        if t:
            lines.append(t)

    # tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                t = norm(cell.text)
                if t:
                    # cell.text can contain multiple lines
                    for part in t.split("\n"):
                        part = norm(part)
                        if part:
                            lines.append(part)

    return lines

def main():
    if not DOCX_PATH.exists():
        raise FileNotFoundError(f"❌ DOCX not found: {DOCX_PATH}")

    doc = Document(str(DOCX_PATH))
    lines = collect_all_text(doc)

    # Merge continuation lines for numbered items
    items = []
    cur_num = None
    cur_text = ""

    for line in lines:
        m = NUM_LINE_PAT.match(line)
        if m:
            if cur_num is not None:
                items.append(cur_text)
            cur_num = int(m.group(1))
            cur_text = m.group(2).strip()
        else:
            if cur_num is not None:
                cur_text = (cur_text + " " + line).strip()

    if cur_num is not None:
        items.append(cur_text)

    rows = []
    for body in items:
        q, a = split_qa(body)
        if q:
            rows.append({"question_text": q, "answer_text": a})

    # If still empty, try direct split for any line containing '?'
    if not rows:
        for line in lines:
            if "?" in line:
                q, a = split_qa(line)
                if q:
                    rows.append({"question_text": q, "answer_text": a})

    df = pd.DataFrame(rows)
    if not df.empty:
        df = df.drop_duplicates(subset=["question_text"]).reset_index(drop=True)

    df.to_csv(OUT_CSV, index=False, encoding="utf-8")
    print(f"✅ Saved QA pairs: {len(df)}")
    print(f"✅ File: {OUT_CSV}")

if __name__ == "__main__":
    main()